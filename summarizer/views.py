import os
import tempfile
import docx
import PyPDF2
import cv2
import yt_dlp
import librosa
import pytesseract
from PIL import Image
import numpy as np

from django.http import JsonResponse, FileResponse
from django.views.decorators.csrf import csrf_exempt
from django.core.files.storage import default_storage
from django.shortcuts import render

# HuggingFace Summarization
from transformers import pipeline
summarizer = pipeline("summarization", model="facebook/bart-large-cnn")

# =========================
# Helper Functions
# =========================
def extract_text_pdf(file_path):
    """Ekstrak teks dari PDF"""
    try:
        text = ""
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ""
        return text.strip()
    except Exception as e:
        print(f"[ERROR] extract_text_pdf: {e}")
        return ""


def extract_text_docx(file_path):
    """Ekstrak teks dari DOCX"""
    try:
        doc = docx.Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    except Exception as e:
        print(f"[ERROR] extract_text_docx: {e}")
        return ""


def extract_text_image(file_path):
    """Ekstrak teks dari gambar via OCR"""
    try:
        img = Image.open(file_path)
        return pytesseract.image_to_string(img)
    except Exception as e:
        print(f"[ERROR] extract_text_image: {e}")
        return ""


def extract_audio_features(file_path):
    """Ambil fitur audio sederhana (MFCC)"""
    try:
        y, sr = librosa.load(file_path, sr=16000)
        mfcc = librosa.feature.mfcc(y=y, sr=sr, n_mfcc=40)
        return " ".join(map(str, np.mean(mfcc.T, axis=0)))
    except Exception as e:
        print(f"[ERROR] extract_audio_features: {e}")
        return ""


def extract_frames_mp4(file_path):
    """Hitung jumlah frame video (dummy transcription)"""
    try:
        cap = cv2.VideoCapture(file_path)
        frames = 0
        while cap.isOpened():
            ret, _ = cap.read()
            if not ret:
                break
            frames += 1
        cap.release()
        return f"Video with {frames} frames processed (dummy transcription)."
    except Exception as e:
        print(f"[ERROR] extract_frames_mp4: {e}")
        return ""


def download_youtube(url):
    """Download video YouTube pakai yt_dlp"""
    try:
        out_path = tempfile.mktemp(suffix=".mp4")
        ydl_opts = {"format": "best", "outtmpl": out_path}
        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            ydl.download([url])
        return out_path
    except Exception as e:
        print(f"[ERROR] download_youtube: {e}")
        return ""


def run_summarizer(text, max_len=200):
    """Ringkas teks menggunakan HuggingFace"""
    if not text.strip():
        return "No content to summarize"
    try:
        summary = summarizer(
            text[:1000],
            max_length=max_len,
            min_length=50,
            do_sample=False
        )
        return summary[0]["summary_text"]
    except Exception as e:
        print(f"[ERROR] run_summarizer: {e}")
        return "Summarization failed"


def save_summary_to_docx(summary):
    """Simpan hasil ringkasan ke DOCX sementara"""
    try:
        filepath = tempfile.mktemp(suffix=".docx")
        doc = docx.Document()
        doc.add_heading("Summarization Result", level=1)
        doc.add_paragraph(summary)
        doc.save(filepath)
        return filepath
    except Exception as e:
        print(f"[ERROR] save_summary_to_docx: {e}")
        return ""


# =========================
# Views
# =========================
def index(request):
    return render(request, "summarizer/index.html")


@csrf_exempt
def summarize_pdf(request):
    if request.method == "POST" and request.FILES.get("file"):
        file_path = default_storage.path(default_storage.save(
            request.FILES["file"].name, request.FILES["file"]
        ))
        text = extract_text_pdf(file_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)


@csrf_exempt
def summarize_docx(request):
    if request.method == "POST" and request.FILES.get("file"):
        file_path = default_storage.path(default_storage.save(
            request.FILES["file"].name, request.FILES["file"]
        ))
        text = extract_text_docx(file_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)


@csrf_exempt
def summarize_image(request):
    if request.method == "POST" and request.FILES.get("file"):
        file_path = default_storage.path(default_storage.save(
            request.FILES["file"].name, request.FILES["file"]
        ))
        text = extract_text_image(file_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)


@csrf_exempt
def summarize_mp4(request):
    if request.method == "POST" and request.FILES.get("file"):
        file_path = default_storage.path(default_storage.save(
            request.FILES["file"].name, request.FILES["file"]
        ))
        text = extract_frames_mp4(file_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)


@csrf_exempt
def summarize_audio(request, fmt=None):
    if request.method == "POST" and request.FILES.get("file"):
        file_path = default_storage.path(default_storage.save(
            request.FILES["file"].name, request.FILES["file"]
        ))
        text = extract_audio_features(file_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)


@csrf_exempt
def summarize_youtube(request):
    if request.method == "POST":
        url = request.POST.get("url")
        if not url:
            return JsonResponse({"error": "No YouTube URL provided"}, status=400)

        video_path = download_youtube(url)
        if not video_path:
            return JsonResponse({"error": "Failed to download video"}, status=500)

        text = extract_frames_mp4(video_path)
        summary = run_summarizer(text)
        docx_file = save_summary_to_docx(summary)
        return FileResponse(open(docx_file, "rb"),
                            as_attachment=True,
                            filename="summary.docx")
    return JsonResponse({"error": "Invalid request"}, status=400)
